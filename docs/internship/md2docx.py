"""Markdown -> DOCX for the BUAD920 internship proposal.

Handles the constructs actually used in the source: ATX headings, pipe tables
(with alignment row), '-' bullets, '1.' ordered lists, '---' rules, and inline
**bold** / *italic* / `code`. Anything else is emitted as body text.
"""
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC, OUT = sys.argv[1], sys.argv[2]

INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)      # deep navy
MUTED = RGBColor(0x55, 0x5555 >> 8 & 0xFF, 0x55)
BODY_FONT = 'Times New Roman'
HEAD_FONT = 'Arial'

doc = Document()

# ---------------------------------------------------------------- page setup
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)      # A4
sec.top_margin = sec.bottom_margin = Cm(2.2)
sec.left_margin = sec.right_margin = Cm(2.2)

# ---------------------------------------------------------------- base styles
normal = doc.styles['Normal']
normal.font.name = BODY_FONT
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal._element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

HEADS = {
    'Heading 1': (16, True, ACCENT, 18, 8),
    'Heading 2': (13, True, ACCENT, 16, 6),
    'Heading 3': (11.5, True, INK, 12, 4),
    'Heading 4': (11, True, INK, 10, 3),
}
for name, (size, bold, colour, before, after) in HEADS.items():
    st = doc.styles[name]
    st.font.name = HEAD_FONT
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = colour
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    st._element.rPr.rFonts.set(qn('w:eastAsia'), HEAD_FONT)


def shade(cell, hexcolour):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:fill'), hexcolour)
    cell._tc.get_or_add_tcPr().append(el)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


INLINE = re.compile(r'(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)')


def emit_runs(par, text, base_bold=False, size=None):
    for piece in INLINE.split(text):
        if not piece:
            continue
        bold, italic, mono = base_bold, False, False
        if piece.startswith('**') and piece.endswith('**'):
            piece, bold = piece[2:-2], True
        elif piece.startswith('*') and piece.endswith('*'):
            piece, italic = piece[1:-1], True
        elif piece.startswith('`') and piece.endswith('`'):
            piece, mono = piece[1:-1], True
        run = par.add_run(piece)
        run.bold = bold
        run.italic = italic
        if mono:
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
        elif size:
            run.font.size = Pt(size)


def horizontal_rule():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'BFC7D1')
    borders.append(bottom)
    # w:pBdr must precede w:spacing / w:ind in the CT_PPr sequence.
    anchor = pPr.find(qn('w:spacing'))
    if anchor is None:
        anchor = pPr.find(qn('w:ind'))
    if anchor is None:
        pPr.append(borders)
    else:
        anchor.addprevious(borders)


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def add_table(header, rows):
    ncols = len(header)
    t = doc.add_table(rows=1, cols=ncols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    hdr = t.rows[0]
    set_repeat_header(hdr)
    for i, text in enumerate(header):
        cell = hdr.cells[i]
        cell.text = ''
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.space_before = Pt(2)
        emit_runs(par, text, base_bold=True, size=9.5)
        for r in par.runs:
            r.font.name = HEAD_FONT
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, '1F3A5F')
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i in range(ncols):
            cell = cells[i]
            cell.text = ''
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.space_before = Pt(2)
            emit_runs(par, row[i] if i < len(row) else '', size=9.5)
            if n % 2 == 1:
                shade(cell, 'F2F5F9')
    doc.add_paragraph().paragraph_format.space_after = Pt(4)



def gather(idx, first):
    """Absorb soft-wrapped continuation lines into one logical block."""
    buf = [first]
    while idx < len(lines):
        nxt = lines[idx].strip()
        if (not nxt or nxt.startswith(('|', '#', '- ', '* ')) or
                re.match(r'^-{3,}$', nxt) or re.match(r'^\d+\.\s', nxt)):
            break
        buf.append(nxt)
        idx += 1
    return ' '.join(buf), idx


lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
first_heading_done = False
while i < len(lines):
    raw = lines[i]
    line = raw.rstrip()
    stripped = line.strip()

    if not stripped:
        i += 1
        continue

    # table
    if stripped.startswith('|') and i + 1 < len(lines) and re.match(
            r'^\|[\s:|-]+\|$', lines[i + 1].strip()):
        header = split_row(stripped)
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append(split_row(lines[i].strip()))
            i += 1
        add_table(header, rows)
        continue

    if re.match(r'^-{3,}$', stripped):
        horizontal_rule()
        i += 1
        continue

    m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
    if m:
        level, text = len(m.group(1)), m.group(2)
        if level == 1 and first_heading_done:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = doc.add_paragraph(style=f'Heading {level}')
        emit_runs(p, text)
        first_heading_done = True
        i += 1
        continue

    m = re.match(r'^[-*]\s+(.*)$', stripped)
    if m:
        text, i = gather(i + 1, m.group(1))
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        emit_runs(p, text)
        continue

    m = re.match(r'^\d+\.\s+(.*)$', stripped)
    if m:
        text, i = gather(i + 1, m.group(1))
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(3)
        emit_runs(p, text)
        continue

    # paragraph: join soft-wrapped continuation lines
    text, i = gather(i + 1, stripped)
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text.startswith('*') and text.endswith('*') and '**' not in text:
        emit_runs(p, text)
    else:
        emit_runs(p, text)

# ------------------------------------------------------------------ footer
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run('BUAD920 Internship Proposal — Shanika Kirillawala (20250356) — '
               'Wizards Learning Hub    |    Page ')
r.font.size = Pt(8)
r.font.name = HEAD_FONT
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
inner = OxmlElement('w:r')
rPr = OxmlElement('w:rPr')
sz = OxmlElement('w:sz')
sz.set(qn('w:val'), '16')
rPr.append(sz)
inner.append(rPr)
t = OxmlElement('w:t')
t.text = '1'
inner.append(t)
fld.append(inner)
fp._p.append(fld)

# python-docx ships a settings.xml whose <w:zoom> omits the required percent.
zoom = doc.settings.element.find(qn('w:zoom'))
if zoom is not None and zoom.get(qn('w:percent')) is None:
    zoom.set(qn('w:percent'), '100')

doc.save(OUT)
print('wrote', OUT)
